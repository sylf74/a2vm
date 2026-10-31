# SPDX-License-Identifier: LicenseRef-a2vm-proprietary
"""Documents of the conversational search (LLD triage-prioritization §16.6): Word, PowerPoint and
CSV built from the decision service's data, never written by the model.

A document carries the run it comes from, the filters that chose its decisions, the counts per
level, the table of decisions and, for the most urgent ones, the explanation stored with the
decision (rationale, SSVC inputs and their sources, reasons, fixes). Labels are in the client's language
(`run.language`, v0.17.3, from `verdict.i18n`), as the reports of `verdict run`; the texts of the
decisions (rationales, reasons) stay in English. Text from vendors is kept as text: no link is made from it.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt
from pptx import Presentation
from pptx.util import Cm as PCm
from pptx.util import Pt as PPt

from ..i18n import has, lang_of, t

FORMATS = {
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx"),
    "csv": ("text/csv; charset=utf-8", "csv"),
}
LEVELS = ("ACT", "ATTEND", "TRACK")
COLUMNS = ["level", "cve", "asset", "ssvc", "kev", "epss", "cvss", "due_date", "product", "title"]
DOCX_WIDTHS_CM = (1.6, 2.9, 3.4, 1.3, 1.1, 1.4, 1.2, 2.0, 3.0, 8.8)  # 26.7 cm: A4 landscape, 1.5 cm margins
FILTERS = ("level", "cve", "asset", "exposure", "zone", "kev", "needs_review", "source", "min_epss", "min_cvss")
FILTERS += ("product", "q")
SSVC_INPUTS = ("exploitation", "automatable", "technical_impact")


def _label(lang: str, key: str) -> str:
    return t(lang, f"x.col.{key}")


def _bod(lang: str, code: str) -> str:
    return t(lang, f"bod.{code}") if has(f"bod.{code}") else code


def cell(item: dict[str, Any], key: str, lang: str = "en") -> str:
    value = item.get(key)
    if value is None or value == "":
        return "—"
    if isinstance(value, bool):
        return t(lang, "yes") if value else t(lang, "no")
    if key == "epss" and isinstance(value, (int, float)):
        return f"{value:.1%}"
    return str(value)


def filters_text(filters: dict[str, str], lang: str = "en") -> str:
    parts = [f"{t(lang, 'x.f.' + k)} = {v}" for k, v in filters.items() if k in FILTERS]
    return t(lang, "x.sep").join(parts) or t(lang, "x.filters.none")


def _input(detail: dict[str, Any], key: str, lang: str) -> str:
    v = detail.get(key) or {}
    if isinstance(v, dict):
        return t(lang, "x.source", value=v.get("value", "—"), source=v.get("source", "—"))
    return str(v)


def _explanation(detail: dict[str, Any], lang: str = "en") -> list[tuple[str, str]]:
    sep = t(lang, "x.sep")
    lines = [(t(lang, "x.rationale"), detail.get("rationale") or "—")]
    lines += [(t(lang, f"x.{key}"), _input(detail, key, lang)) for key in SSVC_INPUTS]
    exposure = t(lang, "x.source", value=detail.get("exposure", "—"), source=detail.get("exposure_source", "—"))
    lines.append((t(lang, "x.exposure"), exposure))
    if detail.get("reasons"):
        lines.append((t(lang, "x.reasons"), sep.join(detail["reasons"])))
    if detail.get("review_reasons"):
        lines.append((t(lang, "x.review"), sep.join(detail["review_reasons"])))
    fixes = detail.get("fixes") or []
    more = f" (+{len(fixes) - 10})" if len(fixes) > 10 else ""
    lines.append((t(lang, "x.fixes"), sep.join(fixes[:10]) + more if fixes else t(lang, "x.fixes.none")))
    if detail.get("rationale_source"):
        lines.append((t(lang, "x.rationale_source"), detail["rationale_source"]))
    return lines


def build(kind: str, meta: dict[str, Any], items: list[dict[str, Any]], details: list[dict[str, Any]]) -> bytes:
    """meta: tenant, run {id, generated_at, language}, filters, total, counts {level: n}, truncated;
    meta["language"], when set, wins over the run's."""
    lang = lang_of(meta.get("language") or (meta.get("run") or {}).get("language"))
    if kind == "csv":
        return _csv(items, lang)
    if kind == "docx":
        return _docx(meta, items, details, lang)
    if kind == "pptx":
        return _pptx(meta, items, details, lang)
    raise ValueError(f"format: one of {', '.join(FORMATS)}")


def _csv(items: list[dict[str, Any]], lang: str = "en") -> bytes:
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["id", *(_label(lang, k) for k in COLUMNS)])
    for it in items:
        w.writerow([it.get("id", ""), *(cell(it, k, lang) for k in COLUMNS)])
    return ("﻿" + out.getvalue()).encode("utf-8")  # BOM: Excel reads the accents


def _header_lines(meta: dict[str, Any], lang: str = "en") -> list[str]:
    counts = meta.get("counts") or {}
    at = meta["run"]["generated_at"][:16].replace("T", " ")
    lines = [
        t(lang, "x.run", run=meta["run"]["id"], at=at),
        t(lang, "x.filters", filters=filters_text(meta.get("filters") or {}, lang)),
        t(lang, "x.total", n=meta["total"], counts=", ".join(f"{lv} {counts.get(lv, 0)}" for lv in LEVELS)),
    ]
    if meta.get("truncated"):
        lines.append(t(lang, "x.limited", n=meta["shown"]))
    lines.append(t(lang, "x.method"))
    return lines


def _docx(meta: dict[str, Any], items: list[dict[str, Any]], details: list[dict[str, Any]], lang: str = "en") -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)  # A4 landscape
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))
    doc.styles["Normal"].font.size = Pt(9)
    doc.add_heading(t(lang, "x.title", tenant=meta["tenant"]), level=1)
    for line in _header_lines(meta, lang):
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=len(COLUMNS))
    table.style = "Table Grid"
    table.autofit = False
    for i, key in enumerate(COLUMNS):
        table.rows[0].cells[i].text = _label(lang, key)
        for r in table.rows[0].cells[i].paragraphs[0].runs:
            r.bold = True
    for it in items:
        cells = table.add_row().cells
        for i, key in enumerate(COLUMNS):
            cells[i].text = cell(it, key, lang)
    for i, width in enumerate(DOCX_WIDTHS_CM):  # the grid (LibreOffice) and every cell (Word)
        table.columns[i].width = Cm(width)
        for c in table.columns[i].cells:
            c.width = Cm(width)
    if details:
        doc.add_heading(t(lang, "x.explain"), level=2)
        for d in details:
            asset = d.get("asset_id_full") or d.get("asset_id", "")
            on = t(lang, "x.on", level=d.get("delivery_decision", ""), cve=d.get("cve_id", ""), asset=asset)
            doc.add_heading(on, level=3)
            if d.get("due_date"):
                doc.add_paragraph(t(lang, "x.due", date=d["due_date"], delay=_bod(lang, d.get("bod_timeline", ""))))
            for label, text in _explanation(d, lang):
                p = doc.add_paragraph()
                p.add_run(label + t(lang, "x.colon")).bold = True
                p.add_run(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _pptx(meta: dict[str, Any], items: list[dict[str, Any]], details: list[dict[str, Any]], lang: str = "en") -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = PCm(33.867), PCm(19.05)  # 16:9
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = f"A2VM — {meta['tenant']}"
    title.placeholders[1].text = "\n".join(_header_lines(meta, lang))
    for p in title.placeholders[1].text_frame.paragraphs:
        for r in p.runs:
            r.font.size = PPt(14)
    cols = [c for c in COLUMNS if c != "title"]
    per_slide = 12
    for start in range(0, len(items), per_slide):
        chunk = items[start : start + per_slide]
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = t(lang, "x.slides", start=start + 1, end=start + len(chunk), total=meta["total"])
        shape = slide.shapes.add_table(
            len(chunk) + 1, len(cols), PCm(1), PCm(3.2), PCm(31.8), PCm(1) * (len(chunk) + 1)
        )
        tbl = shape.table
        for i, key in enumerate(cols):
            tbl.cell(0, i).text = _label(lang, key)
        for r, it in enumerate(chunk, 1):
            for i, key in enumerate(cols):
                tbl.cell(r, i).text = cell(it, key, lang)
        for row in tbl.rows:
            for c in row.cells:
                for p in c.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = PPt(10)
    for d in details:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"{d.get('delivery_decision', '')} — {d.get('cve_id', '')} — {d.get('asset_id', '')}"
        body = slide.placeholders[1].text_frame
        body.clear()
        first = True
        for label, text in _explanation(d, lang):
            p = body.paragraphs[0] if first else body.add_paragraph()
            first = False
            p.text = label + t(lang, "x.colon") + text[:400]
            for run in p.runs:
                run.font.size = PPt(12)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
