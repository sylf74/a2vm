# SPDX-License-Identifier: LicenseRef-a2vm-proprietary
"""Build the neutral Word templates of the alert sheet, one per language of the catalogue
(via/src/via/alert/templates/alert.<lang>.docx: en, the reference, and fr).

No client name, logo or classification: a client's own template goes in its project
(templates/alert.<lang>.docx, or templates/alert.docx for every language). Run again after changing
this script; the result is committed.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

DIR = Path(__file__).resolve().parents[1] / "via" / "src" / "via" / "alert" / "templates"
TEXT = {
    "en": {
        "sheet_title": "Alert sheet",
        "meta": "Level {{ niveau }} — {{ date }}{% if source %} — source: {{ source }}{% endif %}",
        "summary": "Summary",
        "score_open": " [CVSS v{{ c.cvss_version }} score: ",
        "cve_link": "CVE link: {{r c.lien_cve }}",
        "cvss_link": "CVSS link: {{r c.lien_cvss }}",
        "poc": (
            "{% if poc %}A proof of concept or an exploit is available in open sources."
            "{% else %}No proof of concept is available in open sources to date.{% endif %}"
        ),
        "why": "Why this alert",
        "assets": "Your assets concerned",
        "assets_none": "No asset identified with certainty.",
        "info": "Information",
        "risks": "Risks",
        "risks_none": "Not specified by the sources.",
        "criticality": "Severity",
        "score": "CVSS v{{ cvss_version }} score: ",
        "impact": "Main impact",
        "exploited": "The vulnerability is actively exploited",
        "fix_exists": "A fix exists",
        "workaround_exists": "A temporary workaround exists",
        "type": "Type of the vulnerability",
        "cwe_line": "{{ w.id }}{% if w.nom %}: {{ w.nom }}{% endif %}",
        "details": "Exploitation details",
        "regarding": "Regarding ",
        "av": "Attack vector: {{ c.vecteur }}",
        "ac": "Attack complexity: {{ c.complexite }}",
        "pr": "Privileges required for the attack: {{ c.privileges }}",
        "ui": "The attacker needs a user's interaction: {{ c.interaction }}",
        "priv": "Exploiting the vulnerability gives privileged rights: {{ c.droits_privilegies }}",
        "components": "Vulnerable components",
        "recommendations": "Recommendations",
        "delay": "Recommended remediation delay: {{ delai }} (CISA BOD 26-04).",
        "fixes": "Security fixes to apply",
        "fixes_none": "No fix found in the sources to date.",
        "workarounds": "Workaround",
        "workarounds_none": "No workaround found in the sources to date.",
        "references": "References",
        "table": ["Asset", "Zone", "Exposure", "Installed version", "Fixed version", "Owner", "Source"],
    },
    "fr": {
        "sheet_title": "Fiche d'alerte",
        "meta": "Niveau {{ niveau }} — {{ date }}{% if source %} — source : {{ source }}{% endif %}",
        "summary": "Résumé",
        "score_open": " [Score CVSS v{{ c.cvss_version }} : ",
        "cve_link": "Lien CVE : {{r c.lien_cve }}",
        "cvss_link": "Lien CVSS : {{r c.lien_cvss }}",
        "poc": (
            "{% if poc %}Une preuve de concept ou un exploit est disponible en sources ouvertes."
            "{% else %}Actuellement, aucune preuve de concept n’est disponible en sources ouvertes.{% endif %}"
        ),
        "why": "Pourquoi cette alerte",
        "assets": "Vos actifs concernés",
        "assets_none": "Aucun actif identifié avec certitude.",
        "info": "Informations",
        "risks": "Risques",
        "risks_none": "Non précisés par les sources.",
        "criticality": "Criticité",
        "score": "Score CVSS v{{ cvss_version }} : ",
        "impact": "Impact principal",
        "exploited": "La faille est activement exploitée",
        "fix_exists": "Un correctif existe",
        "workaround_exists": "Une mesure de contournement temporaire existe",
        "type": "La vulnérabilité exploitée est de type",
        "cwe_line": "{{ w.id }}{% if w.nom %} : {{ w.nom }}{% endif %}",
        "details": "Détails sur l’exploitation",
        "regarding": "Concernant la ",
        "av": "Vecteur d’attaque : {{ c.vecteur }}",
        "ac": "Complexité de l’attaque : {{ c.complexite }}",
        "pr": "Privilèges nécessaires pour réaliser l’attaque : {{ c.privileges }}",
        "ui": "L’attaquant doit interagir avec la connexion d’un utilisateur : {{ c.interaction }}",
        "priv": "L’exploitation de la faille permet d’obtenir des droits privilégiés : {{ c.droits_privilegies }}",
        "components": "Composants vulnérables",
        "recommendations": "Recommandations",
        "delay": "Délai de remédiation recommandé : {{ delai }} (référentiel CISA BOD 26-04).",
        "fixes": "Mise en place de correctifs de sécurité",
        "fixes_none": "Aucun correctif identifié dans les sources à ce jour.",
        "workarounds": "Solution de contournement",
        "workarounds_none": "Aucune mesure de contournement identifiée dans les sources à ce jour.",
        "references": "Références",
        "table": ["Actif", "Zone", "Exposition", "Version installée", "Version corrigée", "Responsable", "Source"],
    },
}


def main() -> int:
    for lang in TEXT:
        build(lang)
    return 0


def build(lang: str) -> None:
    L = TEXT[lang]  # noqa: N806 - the labels of the language
    out = DIR / f"alert.{lang}.docx"
    doc = Document()
    for name in ("Normal", "List Bullet"):
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.size = Pt(10.5)
    props = doc.core_properties
    props.author, props.title = "A2VM", L["sheet_title"]
    props.comments = "Neutral A2VM template; tags: docs/alert-word.md"

    def p(text: str = "", style: str | None = None, bold: bool = False):
        para = doc.add_paragraph(style=style)
        if text:
            run = para.add_run(text)
            run.bold = bold
        return para

    def ctl(tag: str):
        doc.add_paragraph(tag)

    doc.add_heading("{{ titre }}", level=0)
    meta = p(L["meta"])
    meta.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_heading(L["summary"], level=1)
    ctl("{%p for c in cves %}")
    para = p()
    para.add_run("{{ c.id }}").bold = True
    para.add_run(L["score_open"])
    para.add_run("{{ c.cvss_score }}").bold = True
    para.add_run("]")
    p(L["cve_link"])
    ctl("{%p if c.lien_cvss %}")
    p(L["cvss_link"])
    ctl("{%p endif %}")
    ctl("{%p endfor %}")
    ctl("{%p if produit_description %}")
    p("{{ produit_description }}")
    ctl("{%p endif %}")
    p(L["poc"])

    doc.add_heading(L["why"], level=1)
    ctl("{%p for r in pourquoi %}")
    p("{{ r }}", style="List Bullet")
    ctl("{%p endfor %}")

    doc.add_heading(L["assets"], level=1)
    ctl("{%p if actifs %}")
    table = doc.add_table(rows=4, cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    heads = L["table"]
    for cell, text in zip(table.rows[0].cells, heads, strict=True):
        cell.paragraphs[0].add_run(text).bold = True
    table.rows[1].cells[0].paragraphs[0].add_run("{%tr for a in actifs %}")
    fields = ["nom", "zone", "exposition", "version_installee", "version_corrigee", "responsable", "source"]
    for cell, key in zip(table.rows[2].cells, fields, strict=True):
        cell.paragraphs[0].add_run("{{ a." + key + " }}")
    table.rows[3].cells[0].paragraphs[0].add_run("{%tr endfor %}")
    ctl("{%p else %}")
    p(L["assets_none"])
    ctl("{%p endif %}")

    doc.add_heading(L["info"], level=1)
    doc.add_heading(L["risks"], level=2)
    ctl("{%p for r in risques %}")
    p("{{ r }}")
    ctl("{%p else %}")
    p(L["risks_none"])
    ctl("{%p endfor %}")
    doc.add_heading(L["criticality"], level=2)
    para = p(L["score"])
    para.add_run("{{ cvss_score }}").bold = True
    doc.add_heading(L["impact"], level=2)
    p("{{ impact }}")
    doc.add_heading(L["exploited"], level=2)
    p("{{ exploitee }}", style="List Bullet")
    doc.add_heading(L["fix_exists"], level=2)
    p("{{ correctif_existe }}", style="List Bullet")
    doc.add_heading(L["workaround_exists"], level=2)
    p("{{ contournement_existe }}", style="List Bullet")
    doc.add_heading(L["type"], level=2)
    ctl("{%p for w in cwes %}")
    p(L["cwe_line"])
    p("CWE : {{r w.lien }}")
    ctl("{%p endfor %}")
    doc.add_heading(L["details"], level=2)
    ctl("{%p for c in cves %}")
    ctl("{%p if cves|length > 1 %}")
    para = p(L["regarding"])
    para.add_run("{{ c.id }}").bold = True
    para.add_run(" :")
    ctl("{%p endif %}")
    for line in (
        L["av"],
        L["ac"],
        L["pr"],
        L["ui"],
        L["priv"],
    ):
        p(line, style="List Bullet")
    ctl("{%p endfor %}")
    doc.add_heading(L["components"], level=2)
    ctl("{%p for c in cves %}")
    para = p(L["regarding"])
    para.add_run("{{ c.id }}").bold = True
    para.add_run(" :")
    ctl("{%p for v in c.composants %}")
    p("{{ v }}", style="List Bullet")
    ctl("{%p endfor %}")
    ctl("{%p endfor %}")

    doc.add_heading("Description", level=1)
    p("{{ description }}")

    doc.add_heading(L["recommendations"], level=1)
    ctl("{%p if delai %}")
    p(L["delay"])
    ctl("{%p endif %}")
    doc.add_heading(L["fixes"], level=2)
    ctl("{%p for x in correctifs %}")
    p("{{ x }}", style="List Bullet")
    ctl("{%p else %}")
    p(L["fixes_none"])
    ctl("{%p endfor %}")
    doc.add_heading(L["workarounds"], level=2)
    ctl("{%p for x in contournements %}")
    p("{{ x }}", style="List Bullet")
    ctl("{%p else %}")
    p(L["workarounds_none"])
    ctl("{%p endfor %}")

    doc.add_heading(L["references"], level=1)
    ctl("{%p for l in references %}")
    p("{{r l }}")
    ctl("{%p endfor %}")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == "__main__":
    sys.exit(main())
