# SPDX-License-Identifier: LicenseRef-a2vm-proprietary
"""v0.15.1: assets named as people name them, documents of the chat (Word, PowerPoint, CSV) built
from the data and downloaded through the chat's address, retention of the chats."""

from __future__ import annotations

import csv
import io
import json
import logging
import sqlite3
import threading
import time
from pathlib import Path

import httpx
import pytest
import yaml
from docx import Document
from pptx import Presentation

from a2vm.service import Problem
from verdict.service.retention import purge
from verdict.service.tools import serve as serve_tools

from .test_v0150 import CHAT, tools  # noqa: F401 - fixture
from .test_verdict_service import READ, _push, _service, triage  # noqa: F401 - fixture

ROOT = Path(__file__).resolve().parents[1]
CHART = ROOT / "deploy" / "charts" / "a2vm-tenant"


# -- assets ----------------------------------------------------------------------------------------
def test_an_asset_is_found_by_its_id_its_host_name_or_its_short_name(triage, tmp_path):  # noqa: F811
    run, path = triage
    service = _service(tmp_path, run.client.id)
    _push(service, path)
    row = run.rows[0]
    full = row.asset_id_full or row.asset_id
    names = {full, full.split(":", 1)[-1], row.asset_id, full.upper()}
    for name in names:
        page = service.read(run.client.id, READ, "decisions", {"asset": name, "limit": "200"})
        assert page["total"] >= 1 and {i["asset"] for i in page["items"]} == {full}, name
        assert service.read(run.client.id, READ, f"assets/{name}", {})["asset"] == full, name


def test_the_tools_decode_the_path_parameters_the_chat_encodes(tools):  # noqa: F811
    run, t = tools
    full = run.rows[0].asset_id_full or run.rows[0].asset_id
    encoded = full.replace(":", "%3A").replace("/", "%2F")
    assert t.call(CHAT, f"/assets/{encoded}", {})["asset"] == full
    with pytest.raises(Problem) as p:
        t.call(CHAT, "/assets/..%2F..%2Fetc", {})
    assert p.value.status == 400


# -- documents ---------------------------------------------------------------------------------------
@pytest.fixture
def exporting(tools):  # noqa: F811
    run, t = tools
    t.public_url = "https://chat-lab.apps.example.org"
    return run, t


def _fetch(t, link: str) -> tuple[int, dict, bytes]:
    httpd = serve_tools(t, "127.0.0.1", 0)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    try:
        path = link.split("example.org", 1)[1]
        r = httpx.get(f"http://127.0.0.1:{httpd.server_address[1]}{path}")
        return r.status_code, dict(r.headers), r.content
    finally:
        httpd.shutdown()
        httpd.server_close()


def test_word_powerpoint_and_csv_are_built_from_the_data(exporting, caplog):
    run, t = exporting
    caplog.set_level(logging.INFO)
    first_cve = (
        next(r.cve_id for r in run.rows if r.delivery_decision == "ACT")
        if any(r.delivery_decision == "ACT" for r in run.rows)
        else run.rows[0].cve_id
    )
    for kind in ("docx", "pptx", "csv"):
        answer = t.call(CHAT, "/exports", {"format": kind, "explain": "3"})
        link = answer["download"]
        assert link.startswith("https://chat-lab.apps.example.org/a2vm/exports/") and link.endswith(f".{kind}")
        assert answer["decisions"] == min(len(run.rows), 200) and answer["run"]["id"]
        status, headers, body = _fetch(t, link)
        assert (
            status == 200 and "attachment" in headers["content-disposition"] and headers["cache-control"] == "no-store"
        )
        if kind == "docx":
            text = "\n".join(p.text for p in Document(io.BytesIO(body)).paragraphs)
            tables = Document(io.BytesIO(body)).tables
            assert answer["run"]["id"] in text and "Rationale" in text
            assert any(first_cve in c.text for c in tables[0].columns[1].cells)
        elif kind == "pptx":
            texts = [
                sh.text_frame.text
                for s in Presentation(io.BytesIO(body)).slides
                for sh in s.shapes
                if sh.has_text_frame
            ]
            assert any(answer["run"]["id"] in x for x in texts)
        else:
            rows = list(csv.reader(io.StringIO(body.decode("utf-8-sig")), delimiter=";"))
            assert rows[0][:3] == ["id", "Level", "CVE"] and len(rows) == answer["decisions"] + 1
    token = link.split("/a2vm/exports/", 1)[1].split("/", 1)[0]
    server = [r for r in caplog.records if r.name.startswith(("verdict", "a2vm"))]
    assert server and all(token not in r.getMessage() for r in server)  # the link is a credential


def test_a_document_link_is_short_lived_and_exact(exporting):
    _, t = exporting
    link = t.call(CHAT, "/exports", {"format": "csv", "level": "ACT,ATTEND"})["download"]
    assert _fetch(t, link.rsplit("/", 1)[0] + "/other-name.csv")[0] == 404
    assert _fetch(t, link.replace("/a2vm/exports/", "/a2vm/exports/x"))[0] == 404
    for token, (_expires, name, ctype, body) in list(t._exports.items()):
        t._exports[token] = (time.time() - 1, name, ctype, body)
    assert _fetch(t, link)[0] == 404


@pytest.mark.parametrize(
    ("query", "status"), [({"format": "pdf"}, 400), ({"max": "9999"}, 400), ({"explain": "99"}, 400)]
)
def test_bad_document_requests(exporting, query, status):
    _, t = exporting
    with pytest.raises(Problem) as p:
        t.call(CHAT, "/exports", query)
    assert int(p.value.status) == status


def test_no_document_without_the_chat_address(tools):  # noqa: F811
    _, t = tools
    with pytest.raises(Problem) as p:
        t.call(CHAT, "/exports", {"format": "csv"})
    assert p.value.status == 503


# -- retention -------------------------------------------------------------------------------------
def _db(path: Path) -> sqlite3.Connection:
    db = sqlite3.connect(path)
    db.executescript(
        """
        CREATE TABLE chat (id TEXT PRIMARY KEY, user_id TEXT, title TEXT, created_at INTEGER, updated_at INTEGER);
        CREATE TABLE chat_message (id TEXT PRIMARY KEY, chat_id TEXT);
        CREATE TABLE shared_chat (id TEXT PRIMARY KEY, chat_id TEXT);
        CREATE TABLE chat_file (id TEXT PRIMARY KEY, chat_id TEXT, file_id TEXT);
        CREATE TABLE automation_run (id TEXT PRIMARY KEY, chat_id TEXT);
        """
    )
    return db


def test_chats_older_than_the_retention_are_deleted_as_open_webui_does(tmp_path):
    path = tmp_path / "webui.db"
    db = _db(path)
    now = time.time()
    for i, age in enumerate((40, 31, 5, 0)):
        cid = f"c{i}"
        db.execute("INSERT INTO chat VALUES (?, 'u', 't', ?, ?)", (cid, int(now - age * 86400), int(now - age * 86400)))
        db.execute("INSERT INTO chat_message VALUES (?, ?)", (f"m{i}", cid))
        db.execute("INSERT INTO shared_chat VALUES (?, ?)", (f"s{i}", cid))
        db.execute("INSERT INTO chat_file VALUES (?, ?, 'f')", (f"f{i}", cid))
        db.execute("INSERT INTO automation_run VALUES (?, ?)", (f"r{i}", cid))
    db.commit()
    assert purge(path, 30, now) == 2
    assert purge(path, 30, now) == 0
    assert [r[0] for r in db.execute("SELECT id FROM chat ORDER BY id")] == ["c2", "c3"]
    for table in ("chat_message", "shared_chat", "chat_file"):
        assert {r[0] for r in db.execute(f"SELECT chat_id FROM {table}")} == {"c2", "c3"}  # noqa: S608
    assert dict(db.execute("SELECT id, chat_id FROM automation_run").fetchall()) == {
        "r0": None,
        "r1": None,
        "r2": "c2",
        "r3": "c3",
    }
    assert purge(path, 0, now) == 0  # 0 keeps everything


def test_retention_stops_rather_than_guess_when_open_webui_changed(tmp_path):
    path = tmp_path / "webui.db"
    db = sqlite3.connect(path)
    db.execute("CREATE TABLE chat (id TEXT PRIMARY KEY, modified INTEGER)")
    db.commit()
    with pytest.raises(RuntimeError, match="retention stopped"):
        purge(path, 30)
    assert purge(tmp_path / "empty.db", 30) == 0  # no table yet: nothing to do


# -- chart and deploy job ----------------------------------------------------------------------------
def test_chart_publishes_the_downloads_only_and_runs_the_retention():
    text = (CHART / "templates" / "chat.yaml").read_text()
    for needle in (
        "- path: /a2vm/exports/",
        "--public-url=",
        "--export-minutes=",
        'command: ["verdict", "chat-retention"]',
        "--days={{ $chat.retentionDays }}",
        "{{- if gt (int $chat.retentionDays) 0 }}",
    ):
        assert needle in text, needle
    ingress = text.split("kind: Ingress", 1)[1].split("---", 1)[0]
    assert ingress.count("- path:") == 2
    values = yaml.safe_load((CHART / "values.yaml").read_text())["chat"]
    assert values["retentionDays"] == 30 and values["exportMinutes"] == 60
    prompt = values["systemPrompt"]
    assert "list_assets" in prompt and "export_decisions" in prompt and "Source : run" in prompt


def test_chat_values_take_the_retention(tmp_path):
    from .test_v0150 import _chat_values

    out = tmp_path / "v.json"
    base = {"A2VM_CHAT_HOST": "c", "A2VM_LLM_BASE_URL": "http://192.0.2.10:8000/v1", "A2VM_LLM_MODEL": "m"}
    assert _chat_values({"OUT": str(out), **base, "A2VM_CHAT_RETENTION_DAYS": "90"}).returncode == 0
    assert json.loads(out.read_text())["chat"]["retentionDays"] == 90
    r = _chat_values({"OUT": str(out), **base, "A2VM_CHAT_RETENTION_DAYS": "-1"})
    assert r.returncode == 1 and "RETENTION" in r.stderr
